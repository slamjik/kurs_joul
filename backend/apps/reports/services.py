"""
Сервисный слой модуля «Экспорт отчётов».
Вся бизнес-логика:
- Выгрузка учебной нагрузки в Excel (.xlsx) с фирменным стилем, шапкой кафедры и автофильтрами
- Выгрузка ведомостей оценок в Excel (.xlsx) с цветовой индикацией оценок и статистикой
- Генерация официальных PDF-документов на базе ReportLab с полной поддержкой кириллицы
- Fallback на печатный HTML при отсутствии генератора PDF
"""

from datetime import datetime
import io
import logging
import os
from typing import Any, Dict, List, Optional, Tuple

from django.http import HttpResponse
from django.template.loader import render_to_string
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from grades.models import Grade
from workload.models import Workload

logger = logging.getLogger(__name__)

# Фирменная палитра КафИС
BRAND_BLUE_HEX = "1A56DB"
BRAND_NAVY_HEX = "1E3A8A"
BG_LIGHT_HEX = "F8FAFC"
BG_SUBTITLE_HEX = "F1F5F9"
BORDER_COLOR_HEX = "CBD5E0"


def _sanitize_cell(val: Any) -> Any:
    """Защита от Formula / CSV Injection (CWE-1236)."""
    if isinstance(val, str) and val.startswith(("=", "+", "-", "@", "\t", "\r")):
        return f"'{val}"
    return val


def _apply_clean_excel_styling(
    ws,
    header_row: int = 4,
    data_start_row: int = 5,
    last_row: int = 5,
    total_cols: int = 9,
) -> None:
    """Применяет аккуратное оформление к таблице Excel: границы, зебра, ширина колонок."""
    thin_border = Border(
        left=Side(style="thin", color=BORDER_COLOR_HEX),
        right=Side(style="thin", color=BORDER_COLOR_HEX),
        top=Side(style="thin", color=BORDER_COLOR_HEX),
        bottom=Side(style="thin", color=BORDER_COLOR_HEX),
    )
    header_fill = PatternFill(start_color=BRAND_BLUE_HEX, end_color=BRAND_BLUE_HEX, fill_type="solid")
    header_font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    zebra_fill = PatternFill(start_color=BG_LIGHT_HEX, end_color=BG_LIGHT_HEX, fill_type="solid")
    white_fill = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")

    # Стилизация шапки
    ws.row_dimensions[header_row].height = 26
    for col_idx in range(1, total_cols + 1):
        cell = ws.cell(row=header_row, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin_border

    # Стилизация строк данных
    for r_idx in range(data_start_row, last_row + 1):
        ws.row_dimensions[r_idx].height = 20
        fill_to_use = zebra_fill if (r_idx % 2 == 0) else white_fill
        for c_idx in range(1, total_cols + 1):
            cell = ws.cell(row=r_idx, column=c_idx)
            cell.border = thin_border
            if not cell.fill or cell.fill.start_color.rgb == "00000000":
                cell.fill = fill_to_use
            if not cell.font.name:
                cell.font = Font(name="Arial", size=9.5)

    # Автоподбор ширины столбцов
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        # Пропускаем первые 3 строки (заголовок/подзаголовок/отступ) при расчете ширины
        for cell in col[3:]:
            val_str = str(cell.value or "")
            if len(val_str) > max_len:
                max_len = len(val_str)
        ws.column_dimensions[col_letter].width = min(max(max_len + 4, 11), 42)

    # Включаем сетку
    if ws.views.sheetView:
        ws.views.sheetView[0].showGridLines = True


def export_workload_excel(params: Dict[str, Any], user=None) -> HttpResponse:
    """
    Экспортирует распределение учебной нагрузки в формате .xlsx с презентабельным стилем.
    """
    qs = Workload.objects.select_related("teacher", "discipline", "group").all()

    # Защита от BOLA (Horizontal Privilege Escalation)
    if user and user.is_authenticated and user.role == "teacher":
        qs = qs.filter(teacher__user=user)
    else:
        teacher_id = params.get("teacher")
        if teacher_id:
            qs = qs.filter(teacher_id=teacher_id)

    semester = params.get("semester")
    if semester:
        qs = qs.filter(semester=semester)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Учебная нагрузка"

    # 1. Заголовочный баннер
    ws.merge_cells("A1:J1")
    ws["A1"] = "КафИС • Новотроицкий филиал НИТУ «МИСИС» • Кафедра ГиСЭН"
    ws["A1"].font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
    ws["A1"].fill = PatternFill(start_color=BRAND_NAVY_HEX, end_color=BRAND_NAVY_HEX, fill_type="solid")
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 26

    # 2. Подзаголовок
    sem_text = f"Семестр: {semester}" if semester else "Все семестры кафедры"
    date_text = datetime.now().strftime("%d.%m.%Y %H:%M")
    ws.merge_cells("A2:J2")
    ws["A2"] = f"Ведомость учета и выполнения учебной нагрузки | {sem_text} | Сформировано: {date_text}"
    ws["A2"].font = Font(name="Arial", size=9, color="475569")
    ws["A2"].fill = PatternFill(start_color=BG_SUBTITLE_HEX, end_color=BG_SUBTITLE_HEX, fill_type="solid")
    ws["A2"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[2].height = 20
    ws.row_dimensions[3].height = 8

    # 3. Шапка таблицы
    headers = [
        "№",
        "Преподаватель",
        "Дисциплина",
        "Группа",
        "Семестр",
        "План (ч)",
        "Факт (ч)",
        "% вып.",
        "Аудитория",
        "Пара / День",
    ]
    ws.append(headers)

    total_plan = 0
    total_fact = 0
    start_row = 5

    for idx, item in enumerate(qs, 1):
        plan = item.hours_plan or 0
        fact = item.hours_fact or 0
        total_plan += plan
        total_fact += fact
        pct = round((fact / plan * 100), 1) if plan > 0 else 0.0

        lesson_info = f"{item.get_day_of_week_display() or ''} ({item.get_lesson_number_display() or ''})".strip()
        if not lesson_info or lesson_info == "()":
            lesson_info = "—"

        ws.append([
            idx,
            _sanitize_cell(item.teacher.full_name),
            _sanitize_cell(item.discipline.name),
            _sanitize_cell(item.group.name),
            _sanitize_cell(item.semester),
            plan,
            fact,
            f"{pct}%",
            _sanitize_cell(item.room or "—"),
            _sanitize_cell(lesson_info),
        ])

    last_data_row = ws.max_row

    # 4. Итоговая строка
    overall_pct = round((total_fact / total_plan * 100), 1) if total_plan > 0 else 0.0
    ws.append([
        "ИТОГО ПО КАФЕДРЕ",
        "",
        "",
        "",
        "",
        total_plan,
        total_fact,
        f"{overall_pct}%",
        "",
        "",
    ])
    total_row_idx = ws.max_row
    ws.merge_cells(f"A{total_row_idx}:E{total_row_idx}")

    total_border = Border(
        top=Side(style="thin", color="1E3A8A"),
        bottom=Side(style="double", color="1E3A8A"),
        left=Side(style="thin", color=BORDER_COLOR_HEX),
        right=Side(style="thin", color=BORDER_COLOR_HEX),
    )
    total_fill = PatternFill(start_color=BG_SUBTITLE_HEX, end_color=BG_SUBTITLE_HEX, fill_type="solid")
    for col_i in range(1, len(headers) + 1):
        c = ws.cell(row=total_row_idx, column=col_i)
        c.border = total_border
        c.fill = total_fill
        c.font = Font(name="Arial", size=10, bold=True, color="1E3A8A")
        if col_i in [6, 7, 8]:
            c.alignment = Alignment(horizontal="center", vertical="center")
        else:
            c.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[total_row_idx].height = 24

    _apply_clean_excel_styling(ws, header_row=4, data_start_row=5, last_row=last_data_row, total_cols=len(headers))

    # Автофильтр
    ws.auto_filter.ref = f"A4:{get_column_letter(len(headers))}{last_data_row}"

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    filename = f"workload_{semester or 'all'}_{datetime.now().strftime('%Y%m%d')}.xlsx"
    response = HttpResponse(
        output.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    return response


def export_grades_excel(params: Dict[str, Any], user=None) -> HttpResponse:
    """
    Экспортирует ведомость успеваемости студентов в формате .xlsx с цветовыми акцентами.
    """
    qs = Grade.objects.select_related("student__group", "discipline", "teacher").all()

    # Защита от BOLA (Horizontal Privilege Escalation)
    if user and user.is_authenticated and user.role == "teacher":
        qs = qs.filter(teacher__user=user)

    semester = params.get("semester")
    group_id = params.get("group")
    discipline_id = params.get("discipline")

    if semester:
        qs = qs.filter(semester=semester)
    if group_id:
        qs = qs.filter(student__group_id=group_id)
    if discipline_id:
        qs = qs.filter(discipline_id=discipline_id)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Ведомость успеваемости"

    # 1. Заголовочный баннер
    ws.merge_cells("A1:I1")
    ws["A1"] = "КафИС • Новотроицкий филиал НИТУ «МИСИС» • Кафедра ГиСЭН"
    ws["A1"].font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
    ws["A1"].fill = PatternFill(start_color=BRAND_NAVY_HEX, end_color=BRAND_NAVY_HEX, fill_type="solid")
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 26

    # 2. Подзаголовок
    sem_text = f"Семестр: {semester}" if semester else "Все семестры"
    date_text = datetime.now().strftime("%d.%m.%Y %H:%M")
    ws.merge_cells("A2:I2")
    ws["A2"] = f"Экзаменационная / аттестационная ведомость успеваемости | {sem_text} | Сформировано: {date_text}"
    ws["A2"].font = Font(name="Arial", size=9, color="475569")
    ws["A2"].fill = PatternFill(start_color=BG_SUBTITLE_HEX, end_color=BG_SUBTITLE_HEX, fill_type="solid")
    ws["A2"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[2].height = 20
    ws.row_dimensions[3].height = 8

    # 3. Шапка таблицы
    headers = [
        "№",
        "ФИО студента",
        "Группа",
        "Дисциплина",
        "Оценка",
        "Словесно",
        "Семестр",
        "Преподаватель",
        "Дата выставления",
    ]
    ws.append(headers)

    grades_list = []
    start_row = 5

    def get_verbal_grade(val: float) -> str:
        if val >= 4.5:
            return "Отлично"
        if val >= 3.5:
            return "Хорошо"
        if val >= 3.0:
            return "Удовлетворительно"
        return "Неудовлетворительно"

    for idx, g in enumerate(qs, 1):
        teacher_name = g.teacher.full_name if g.teacher else "—"
        g_val = float(g.grade)
        grades_list.append(g_val)
        verbal = get_verbal_grade(g_val)
        date_str = g.date.strftime("%d.%m.%Y") if g.date else "—"

        ws.append([
            idx,
            _sanitize_cell(g.student.full_name),
            _sanitize_cell(g.student.group.name),
            _sanitize_cell(g.discipline.name),
            g_val,
            verbal,
            _sanitize_cell(g.semester),
            _sanitize_cell(teacher_name),
            _sanitize_cell(date_str),
        ])

        curr_row = ws.max_row
        grade_cell = ws.cell(row=curr_row, column=5)
        grade_cell.alignment = Alignment(horizontal="center", vertical="center")
        if g_val >= 4.5:
            grade_cell.font = Font(name="Arial", size=10, bold=True, color="15803D")  # Green
        elif g_val < 3.0:
            grade_cell.font = Font(name="Arial", size=10, bold=True, color="DC2626")  # Red

    last_data_row = ws.max_row

    # 4. Итоговая строка
    avg_grade = round(sum(grades_list) / len(grades_list), 2) if grades_list else 0.0
    ws.append([
        f"ИТОГО: {len(grades_list)} оценок",
        "",
        "",
        "СРЕДНИЙ БАЛЛ:",
        avg_grade,
        "",
        "",
        "",
        "",
    ])
    total_row_idx = ws.max_row
    ws.merge_cells(f"A{total_row_idx}:C{total_row_idx}")

    total_border = Border(
        top=Side(style="thin", color="1E3A8A"),
        bottom=Side(style="double", color="1E3A8A"),
        left=Side(style="thin", color=BORDER_COLOR_HEX),
        right=Side(style="thin", color=BORDER_COLOR_HEX),
    )
    total_fill = PatternFill(start_color=BG_SUBTITLE_HEX, end_color=BG_SUBTITLE_HEX, fill_type="solid")
    for col_i in range(1, len(headers) + 1):
        c = ws.cell(row=total_row_idx, column=col_i)
        c.border = total_border
        c.fill = total_fill
        c.font = Font(name="Arial", size=10, bold=True, color="1E3A8A")
        if col_i in [1, 4, 5]:
            c.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[total_row_idx].height = 24

    _apply_clean_excel_styling(ws, header_row=4, data_start_row=5, last_row=last_data_row, total_cols=len(headers))

    ws.auto_filter.ref = f"A4:{get_column_letter(len(headers))}{last_data_row}"

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    filename = f"grades_{semester or 'all'}_{datetime.now().strftime('%Y%m%d')}.xlsx"
    response = HttpResponse(
        output.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    return response


# =========================================================================
# Генератор официальных PDF документов через ReportLab
# =========================================================================

def _get_reportlab_fonts() -> Tuple[str, str]:
    """Регистрирует шрифты с поддержкой кириллицы (TrueType) в ReportLab."""
    try:
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.ttfonts import TTFont

        candidates = [
            ("C:/Windows/Fonts/arial.ttf", "C:/Windows/Fonts/arialbd.ttf"),
            ("C:/Windows/Fonts/calibri.ttf", "C:/Windows/Fonts/calibrib.ttf"),
            ("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"),
            ("/usr/share/fonts/dejavu/DejaVuSans.ttf", "/usr/share/fonts/dejavu/DejaVuSans-Bold.ttf"),
        ]
        for reg_path, bold_path in candidates:
            if os.path.exists(reg_path) and os.path.exists(bold_path):
                pdfmetrics.registerFont(TTFont("KafIS-Regular", reg_path))
                pdfmetrics.registerFont(TTFont("KafIS-Bold", bold_path))
                return "KafIS-Regular", "KafIS-Bold"
    except Exception as e:
        logger.warning(f"Ошибка регистрации системного шрифта для ReportLab: {e}")

    return "Helvetica", "Helvetica-Bold"


class _NumberedCanvas:
    """Двухпроходный Canvas ReportLab для вывода сквозной нумерации страниц 'Стр. X из Y'."""

    @staticmethod
    def create_class(page_orientation="portrait", font_name="Helvetica"):
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import A4, landscape
        from reportlab.pdfgen import canvas

        page_w = landscape(A4)[0] if page_orientation == "landscape" else A4[0]

        class CustomCanvas(canvas.Canvas):
            def __init__(self, *args, **kwargs):
                super().__init__(*args, **kwargs)
                self._saved_page_states = []

            def showPage(self):
                self._saved_page_states.append(dict(self.__dict__))
                self._startPage()

            def save(self):
                num_pages = len(self._saved_page_states)
                for state in self._saved_page_states:
                    self.__dict__.update(state)
                    self.draw_page_decorations(num_pages)
                    super().showPage()
                super().save()

            def draw_page_decorations(self, total_pages):
                self.saveState()
                # Нижний колонтитул
                self.setFillColor(colors.HexColor("#64748B"))
                self.setStrokeColor(colors.HexColor("#E2E8F0"))
                self.setLineWidth(0.5)
                self.line(25, 26, page_w - 25, 26)

                try:
                    self.setFont(font_name, 7.5)
                except Exception:
                    self.setFont("Helvetica", 7.5)

                gen_str = f"КафИС • НФ НИТУ МИСИС • Сформировано {datetime.now().strftime('%d.%m.%Y %H:%M')}"
                page_str = f"Страница {self._pageNumber} из {total_pages}"
                try:
                    self.drawString(25, 14, gen_str)
                    self.drawRightString(page_w - 25, 14, page_str)
                except Exception:
                    self.drawString(25, 14, f"KafIS • MISIS • {datetime.now().strftime('%d.%m.%Y %H:%M')}")
                    self.drawRightString(page_w - 25, 14, f"Page {self._pageNumber} of {total_pages}")
                self.restoreState()

        return CustomCanvas



def _render_pdf_or_fallback(template_name: str, context: Dict[str, Any], filename: str) -> HttpResponse:
    """
    Резервный генератор HTML при недоступности библиотек генерации PDF.
    """
    html_content = render_to_string(template_name, context)
    response = HttpResponse(html_content, content_type="text/html; charset=utf-8")
    response["X-PDF-Fallback"] = "True"
    return response


def export_workload_pdf(params: Dict[str, Any], user=None) -> HttpResponse:
    """
    Генерирует официальную PDF-ведомость распределения учебной нагрузки через ReportLab.
    """
    qs = Workload.objects.select_related("teacher", "discipline", "group").all()

    # Защита от BOLA: преподаватель видит только свои часы
    if user and user.is_authenticated and user.role == "teacher":
        qs = qs.filter(teacher__user=user)
    else:
        teacher_id = params.get("teacher")
        if teacher_id:
            qs = qs.filter(teacher_id=teacher_id)

    semester = params.get("semester")
    if semester:
        qs = qs.filter(semester=semester)

    total_plan = sum(item.hours_plan or 0 for item in qs)
    total_fact = sum(item.hours_fact or 0 for item in qs)

    try:
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import A4, landscape
        from reportlab.lib.styles import ParagraphStyle
        from reportlab.platypus import KeepTogether, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

        font_reg, font_bold = _get_reportlab_fonts()

        buffer = io.BytesIO()
        doc = SimpleDocTemplate(
            buffer,
            pagesize=landscape(A4),
            leftMargin=25,
            rightMargin=25,
            topMargin=25,
            bottomMargin=35,
        )

        elements = []

        # Стили
        head_inst_style = ParagraphStyle(
            "InstHead",
            fontName=font_bold,
            fontSize=7.5,
            leading=9.5,
            alignment=1,
            textColor=colors.HexColor("#475569"),
        )
        sub_inst_style = ParagraphStyle(
            "InstSub",
            fontName=font_reg,
            fontSize=7,
            leading=9,
            alignment=1,
            textColor=colors.HexColor("#64748B"),
        )
        title_style = ParagraphStyle(
            "DocTitle",
            fontName=font_bold,
            fontSize=12,
            leading=15,
            alignment=1,
            textColor=colors.HexColor(f"#{BRAND_NAVY_HEX}"),
        )
        meta_style = ParagraphStyle(
            "DocMeta",
            fontName=font_reg,
            fontSize=8,
            leading=10,
            alignment=1,
            textColor=colors.HexColor("#334155"),
        )
        th_style = ParagraphStyle(
            "TH",
            fontName=font_bold,
            fontSize=7.5,
            leading=9.5,
            alignment=1,
            textColor=colors.white,
        )
        td_style = ParagraphStyle(
            "TD",
            fontName=font_reg,
            fontSize=7.5,
            leading=9.5,
        )
        td_center = ParagraphStyle(
            "TDC",
            fontName=font_reg,
            fontSize=7.5,
            leading=9.5,
            alignment=1,
        )

        # Университетская шапка
        elements.append(Paragraph("МИНИСТЕРСТВО НАУКИ И ВЫСШЕГО ОБРАЗОВАНИЯ РОССИЙСКОЙ ФЕДЕРАЦИИ", sub_inst_style))
        elements.append(Paragraph("НОВОТРОИЦКИЙ ФИЛИАЛ НИТУ «МИСИС»", head_inst_style))
        elements.append(Paragraph("КАФЕДРА ГУМАНИТАРНЫХ И СОЦИАЛЬНО-ЭКОНОМИЧЕСКИХ НАУК (ГиСЭН)", head_inst_style))
        elements.append(Spacer(1, 8))
        elements.append(Paragraph("ВЕДОМОСТЬ РАСПРЕДЕЛЕНИЯ УЧЕБНОЙ НАГРУЗКИ", title_style))
        sem_label = f"Учебный семестр: {semester}" if semester else "Все семестры кафедры"
        elements.append(Paragraph(f"{sem_label} | Дата формирования: {datetime.now().strftime('%d.%m.%Y %H:%M')}", meta_style))
        elements.append(Spacer(1, 10))

        # Сводный блок KPI
        pct_compl = round((total_fact / total_plan * 100), 1) if total_plan > 0 else 0.0
        kpi_data = [
            [
                Paragraph(f"<b>Всего записей:</b> {len(qs)}", td_center),
                Paragraph(f"<b>План кафедры:</b> {total_plan} ч.", td_center),
                Paragraph(f"<b>Фактически выполнено:</b> {total_fact} ч.", td_center),
                Paragraph(f"<b>Процент закрытия:</b> {pct_compl}%", td_center),
            ]
        ]
        kpi_table = Table(kpi_data, colWidths=[180, 200, 200, 212])
        kpi_table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor(f"#{BG_SUBTITLE_HEX}")),
            ("BOX", (0, 0), (-1, -1), 0.5, colors.HexColor(f"#{BORDER_COLOR_HEX}")),
            ("INNERGRID", (0, 0), (-1, -1), 0.5, colors.HexColor(f"#{BORDER_COLOR_HEX}")),
            ("TOPPADDING", (0, 0), (-1, -1), 4),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
        ]))
        elements.append(kpi_table)
        elements.append(Spacer(1, 10))

        # Основная таблица
        table_rows = [[
            Paragraph("№", th_style),
            Paragraph("Преподаватель", th_style),
            Paragraph("Дисциплина", th_style),
            Paragraph("Группа", th_style),
            Paragraph("Сем.", th_style),
            Paragraph("План (ч)", th_style),
            Paragraph("Факт (ч)", th_style),
            Paragraph("% вып.", th_style),
            Paragraph("Ауд.", th_style),
            Paragraph("Пара / День", th_style),
        ]]

        for idx, item in enumerate(qs, 1):
            p = item.hours_plan or 0
            f = item.hours_fact or 0
            row_pct = round((f / p * 100), 1) if p > 0 else 0.0
            day_lesson = f"{item.get_day_of_week_display() or ''} ({item.get_lesson_number_display() or ''})".strip()
            if not day_lesson or day_lesson == "()":
                day_lesson = "—"

            table_rows.append([
                Paragraph(str(idx), td_center),
                Paragraph(item.teacher.full_name, td_style),
                Paragraph(item.discipline.name, td_style),
                Paragraph(item.group.name, td_center),
                Paragraph(item.semester, td_center),
                Paragraph(str(p), td_center),
                Paragraph(str(f), td_center),
                Paragraph(f"{row_pct}%", td_center),
                Paragraph(item.room or "—", td_center),
                Paragraph(day_lesson, td_center),
            ])

        col_widths = [24, 140, 178, 65, 48, 52, 52, 50, 48, 135]
        data_table = Table(table_rows, colWidths=col_widths, repeatRows=1)
        data_table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor(f"#{BRAND_BLUE_HEX}")),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor(f"#{BORDER_COLOR_HEX}")),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor(f"#{BG_LIGHT_HEX}")]),
            ("TOPPADDING", (0, 0), (-1, -1), 3),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
        ]))
        elements.append(data_table)
        elements.append(Spacer(1, 15))

        # Блок подписей
        sig_data = [
            [
                Paragraph("<b>Заведующий кафедрой ГиСЭН:</b>", td_style),
                Paragraph("________________ / М.А. Измайлова /", td_style),
                Paragraph("<b>Ответственный по нагрузке:</b>", td_style),
                Paragraph("________________ / ____________ /", td_style),
            ]
        ]
        sig_table = Table(sig_data, colWidths=[180, 220, 180, 212])
        sig_table.setStyle(TableStyle([
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("TOPPADDING", (0, 0), (-1, -1), 6),
        ]))
        elements.append(KeepTogether([sig_table]))

        CustomCanvas = _NumberedCanvas.create_class("landscape", font_name=font_reg)
        doc.build(elements, canvasmaker=CustomCanvas)

        filename = f"workload_{semester or 'all'}_{datetime.now().strftime('%Y%m%d')}"
        response = HttpResponse(buffer.getvalue(), content_type="application/pdf")
        response["Content-Disposition"] = f'attachment; filename="{filename}.pdf"'
        return response

    except Exception as exc:
        logger.error(f"Ошибка формирования PDF через ReportLab: {exc}", exc_info=True)
        context = {
            "items": list(qs),
            "semester": semester,
            "total_plan": total_plan,
            "total_fact": total_fact,
            "generated_date": datetime.now().strftime("%d.%m.%Y %H:%M"),
        }
        filename = f"workload_{semester or 'all'}_{datetime.now().strftime('%Y%m%d')}"
        return _render_pdf_or_fallback("reports/workload_pdf.html", context, filename)


def export_grades_pdf(params: Dict[str, Any], user=None) -> HttpResponse:
    """
    Генерирует официальную PDF-ведомость успеваемости через ReportLab.
    """
    qs = Grade.objects.select_related("student__group", "discipline", "teacher").all()

    # Защита от BOLA: преподаватель видит только свои оценки
    if user and user.is_authenticated and user.role == "teacher":
        qs = qs.filter(teacher__user=user)

    semester = params.get("semester")
    group_id = params.get("group")
    discipline_id = params.get("discipline")

    group_name = ""
    discipline_name = ""

    if semester:
        qs = qs.filter(semester=semester)
    if group_id:
        qs = qs.filter(student__group_id=group_id)
        group_obj = qs.first()
        if group_obj:
            group_name = group_obj.student.group.name
    if discipline_id:
        qs = qs.filter(discipline_id=discipline_id)
        disc_obj = qs.first()
        if disc_obj:
            discipline_name = disc_obj.discipline.name

    try:
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import A4, landscape
        from reportlab.lib.styles import ParagraphStyle
        from reportlab.platypus import KeepTogether, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

        font_reg, font_bold = _get_reportlab_fonts()

        buffer = io.BytesIO()
        doc = SimpleDocTemplate(
            buffer,
            pagesize=landscape(A4),
            leftMargin=25,
            rightMargin=25,
            topMargin=25,
            bottomMargin=35,
        )

        elements = []

        head_inst_style = ParagraphStyle(
            "InstHead",
            fontName=font_bold,
            fontSize=7.5,
            leading=9.5,
            alignment=1,
            textColor=colors.HexColor("#475569"),
        )
        sub_inst_style = ParagraphStyle(
            "InstSub",
            fontName=font_reg,
            fontSize=7,
            leading=9,
            alignment=1,
            textColor=colors.HexColor("#64748B"),
        )
        title_style = ParagraphStyle(
            "DocTitle",
            fontName=font_bold,
            fontSize=12,
            leading=15,
            alignment=1,
            textColor=colors.HexColor(f"#{BRAND_NAVY_HEX}"),
        )
        meta_style = ParagraphStyle(
            "DocMeta",
            fontName=font_reg,
            fontSize=8,
            leading=10,
            alignment=1,
            textColor=colors.HexColor("#334155"),
        )
        th_style = ParagraphStyle(
            "TH",
            fontName=font_bold,
            fontSize=7.5,
            leading=9.5,
            alignment=1,
            textColor=colors.white,
        )
        td_style = ParagraphStyle(
            "TD",
            fontName=font_reg,
            fontSize=7.5,
            leading=9.5,
        )
        td_center = ParagraphStyle(
            "TDC",
            fontName=font_reg,
            fontSize=7.5,
            leading=9.5,
            alignment=1,
        )

        elements.append(Paragraph("МИНИСТЕРСТВО НАУКИ И ВЫСШЕГО ОБРАЗОВАНИЯ РОССИЙСКОЙ ФЕДЕРАЦИИ", sub_inst_style))
        elements.append(Paragraph("НОВОТРОИЦКИЙ ФИЛИАЛ НИТУ «МИСИС»", head_inst_style))
        elements.append(Paragraph("КАФЕДРА ГУМАНИТАРНЫХ И СОЦИАЛЬНО-ЭКОНОМИЧЕСКИХ НАУК (ГиСЭН)", head_inst_style))
        elements.append(Spacer(1, 8))
        elements.append(Paragraph("ЭКЗАМЕНАЦИОННАЯ / АТТЕСТАЦИОННАЯ ВЕДОМОСТЬ", title_style))

        meta_parts = []
        if semester:
            meta_parts.append(f"Семестр: {semester}")
        if group_name:
            meta_parts.append(f"Группа: {group_name}")
        if discipline_name:
            meta_parts.append(f"Дисциплина: {discipline_name}")
        meta_parts.append(f"Дата: {datetime.now().strftime('%d.%m.%Y %H:%M')}")
        elements.append(Paragraph(" | ".join(meta_parts), meta_style))
        elements.append(Spacer(1, 10))

        # Подсчет статистики успеваемости
        grades_vals = [float(g.grade) for g in qs]
        avg_g = round(sum(grades_vals) / len(grades_vals), 2) if grades_vals else 0.0
        excellent_count = sum(1 for g in grades_vals if g >= 4.5)
        risk_count = sum(1 for g in grades_vals if g < 3.0)

        kpi_data = [
            [
                Paragraph(f"<b>Всего оценок:</b> {len(grades_vals)}", td_center),
                Paragraph(f"<b>Средний балл:</b> {avg_g}", td_center),
                Paragraph(f"<b>Отличников:</b> {excellent_count}", td_center),
                Paragraph(f"<b>В зоне риска (&lt; 3.0):</b> {risk_count}", td_center),
            ]
        ]
        kpi_table = Table(kpi_data, colWidths=[180, 200, 200, 212])
        kpi_table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor(f"#{BG_SUBTITLE_HEX}")),
            ("BOX", (0, 0), (-1, -1), 0.5, colors.HexColor(f"#{BORDER_COLOR_HEX}")),
            ("INNERGRID", (0, 0), (-1, -1), 0.5, colors.HexColor(f"#{BORDER_COLOR_HEX}")),
            ("TOPPADDING", (0, 0), (-1, -1), 4),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
        ]))
        elements.append(kpi_table)
        elements.append(Spacer(1, 10))

        table_rows = [[
            Paragraph("№", th_style),
            Paragraph("ФИО студента", th_style),
            Paragraph("Группа", th_style),
            Paragraph("Дисциплина", th_style),
            Paragraph("Оценка", th_style),
            Paragraph("Словесно", th_style),
            Paragraph("Сем.", th_style),
            Paragraph("Преподаватель", th_style),
            Paragraph("Дата", th_style),
        ]]

        def get_verbal_label(val: float) -> str:
            if val >= 4.5:
                return "Отлично"
            if val >= 3.5:
                return "Хорошо"
            if val >= 3.0:
                return "Удовл."
            return "Неуд."

        for idx, g in enumerate(qs, 1):
            g_val = float(g.grade)
            teacher_name = g.teacher.full_name if g.teacher else "—"
            date_str = g.date.strftime("%d.%m.%Y") if g.date else "—"

            g_color = "#15803D" if g_val >= 4.5 else "#DC2626" if g_val < 3.0 else "#1E293B"
            g_style = ParagraphStyle(
                f"Grade_{idx}",
                fontName=font_bold,
                fontSize=8,
                alignment=1,
                textColor=colors.HexColor(g_color),
            )

            table_rows.append([
                Paragraph(str(idx), td_center),
                Paragraph(g.student.full_name, td_style),
                Paragraph(g.student.group.name, td_center),
                Paragraph(g.discipline.name, td_style),
                Paragraph(str(g_val), g_style),
                Paragraph(get_verbal_label(g_val), td_center),
                Paragraph(g.semester, td_center),
                Paragraph(teacher_name, td_style),
                Paragraph(date_str, td_center),
            ])

        col_widths = [24, 150, 65, 175, 45, 55, 48, 155, 75]
        data_table = Table(table_rows, colWidths=col_widths, repeatRows=1)
        data_table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor(f"#{BRAND_BLUE_HEX}")),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor(f"#{BORDER_COLOR_HEX}")),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor(f"#{BG_LIGHT_HEX}")]),
            ("TOPPADDING", (0, 0), (-1, -1), 3),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
        ]))
        elements.append(data_table)
        elements.append(Spacer(1, 15))

        sig_data = [
            [
                Paragraph("<b>Экзаменатор (преподаватель):</b>", td_style),
                Paragraph("________________ / ____________ /", td_style),
                Paragraph("<b>Заведующий кафедрой ГиСЭН:</b>", td_style),
                Paragraph("________________ / М.А. Измайлова /", td_style),
            ]
        ]
        sig_table = Table(sig_data, colWidths=[180, 220, 180, 212])
        sig_table.setStyle(TableStyle([
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("TOPPADDING", (0, 0), (-1, -1), 6),
        ]))
        elements.append(KeepTogether([sig_table]))

        CustomCanvas = _NumberedCanvas.create_class("landscape", font_name=font_reg)
        doc.build(elements, canvasmaker=CustomCanvas)

        filename = f"grades_{semester or 'all'}_{datetime.now().strftime('%Y%m%d')}"
        response = HttpResponse(buffer.getvalue(), content_type="application/pdf")
        response["Content-Disposition"] = f'attachment; filename="{filename}.pdf"'
        return response

    except Exception as exc:
        logger.error(f"Ошибка формирования PDF ведомости через ReportLab: {exc}", exc_info=True)
        context = {
            "items": list(qs),
            "semester": semester,
            "group_name": group_name,
            "discipline_name": discipline_name,
            "generated_date": datetime.now().strftime("%d.%m.%Y %H:%M"),
        }
        filename = f"grades_{semester or 'all'}_{datetime.now().strftime('%Y%m%d')}"
        return _render_pdf_or_fallback("reports/grades_pdf.html", context, filename)
