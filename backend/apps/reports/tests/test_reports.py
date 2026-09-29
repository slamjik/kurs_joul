"""
Тесты модуля «Экспорт отчётов»:
- Экспорт нагрузки в Excel (.xlsx)
- Экспорт оценок в Excel (.xlsx)
- Экспорт нагрузки и оценок в PDF
- Проверка доступности API эндпоинтов и прав
"""

import io
import openpyxl
import pytest
from rest_framework import status
from rest_framework.test import APIClient

from auth_app.services import generate_jwt_tokens_for_user
from auth_app.tests.factories import UserFactory
from grades.tests.factories import GradeFactory
from workload.tests.factories import WorkloadFactory


@pytest.mark.django_db
class TestReportsExport:
    """Тестирование экспорта отчетов в Excel и PDF."""

    def setup_method(self):
        self.client = APIClient()
        self.head_user = UserFactory(role="head")
        self.head_tokens = generate_jwt_tokens_for_user(self.head_user)

    def test_export_workload_excel_success(self):
        WorkloadFactory(semester="2024-1", hours_plan=72, hours_fact=36)

        self.client.credentials(HTTP_AUTHORIZATION=f"Bearer {self.head_tokens['access']}")
        response = self.client.get("/api/reports/workload-excel/?semester=2024-1")

        assert response.status_code == status.HTTP_200_OK
        assert "spreadsheetml.sheet" in response["Content-Type"]

        # Проверяем валидность сформированного Excel файла
        wb = openpyxl.load_workbook(io.BytesIO(response.content))
        ws = wb.active
        assert ws.title == "Учебная нагрузка"
        assert ws.max_row >= 2  # Заголовок + строка данных

    def test_export_grades_excel_success(self):
        GradeFactory(semester="2024-1", grade=4.5)

        self.client.credentials(HTTP_AUTHORIZATION=f"Bearer {self.head_tokens['access']}")
        response = self.client.get("/api/reports/grades-excel/?semester=2024-1")

        assert response.status_code == status.HTTP_200_OK
        assert "spreadsheetml.sheet" in response["Content-Type"]

        wb = openpyxl.load_workbook(io.BytesIO(response.content))
        ws = wb.active
        assert ws.title == "Ведомость успеваемости"
        assert ws.max_row >= 2

    def test_export_workload_pdf_endpoint(self):
        WorkloadFactory(semester="2024-1")

        self.client.credentials(HTTP_AUTHORIZATION=f"Bearer {self.head_tokens['access']}")
        response = self.client.get("/api/reports/workload-pdf/?semester=2024-1")

        assert response.status_code == status.HTTP_200_OK
        # Возвращает либо сгенерированный PDF, либо HTML с печатными стилями при отсутствии GTK
        assert "pdf" in response["Content-Type"] or "html" in response["Content-Type"]

    def test_export_grades_pdf_endpoint(self):
        GradeFactory(semester="2024-1", grade=5.0)

        self.client.credentials(HTTP_AUTHORIZATION=f"Bearer {self.head_tokens['access']}")
        response = self.client.get("/api/reports/grades-pdf/?semester=2024-1")

        assert response.status_code == status.HTTP_200_OK
        assert "pdf" in response["Content-Type"] or "html" in response["Content-Type"]

    def test_workload_export_alias_url(self):
        # Проверяем прямой алиас /api/workload/export/
        self.client.credentials(HTTP_AUTHORIZATION=f"Bearer {self.head_tokens['access']}")
        response = self.client.get("/api/workload/export/")

        assert response.status_code == status.HTTP_200_OK
        assert "spreadsheetml.sheet" in response["Content-Type"]
