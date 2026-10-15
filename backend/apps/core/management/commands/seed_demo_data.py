"""
Management command to seed realistic university data from real semester workload Excel files
and official department roster from Novotroitsk branch of NUST MISIS.

Usage: python manage.py seed_demo_data
"""

import os
import re
import uuid
from datetime import date
from pathlib import Path
import openpyxl

from django.core.management.base import BaseCommand
from django.db import transaction
from django.core.cache import cache

from auth_app.models import User
from workload.models import Department, Teacher, StudyGroup, Discipline, Workload
from grades.models import Student, Grade
from audit.models import AuditLog
from surveys.models import (
    SurveyTemplate,
    SurveyQuestion,
    SurveyAssignment,
    SurveyAnswer,
    TeacherRecommendation,
)


def clean_teacher_name(name):
    if not name:
        return ""
    name = str(name).strip().lstrip("\\/").strip()
    name = re.sub(r"\s+", " ", name)
    # Fix Latin letters in Cyrillic initials
    name = name.replace("C.", "С.").replace("a.", "А.").replace("c.", "с.")
    if name.endswith("А.Н"):
        name = name + "."
    if name.endswith("И.М"):
        name = name + "."
    return name


def parse_hours(val):
    if val is None:
        return 0
    s = str(val).strip()
    if not s:
        return 0
    # Handle formats like "34*2" or "17+17" or "34"
    if "*" in s:
        parts = s.split("*")
        try:
            return int(float(parts[0].strip())) * int(float(parts[1].strip()))
        except Exception:
            pass
    if "+" in s:
        parts = s.split("+")
        try:
            return sum(int(float(p.strip())) for p in parts if p.strip().isdigit())
        except Exception:
            pass
    # Extract leading digits
    m = re.match(r"^(\d+)", s)
    if m:
        try:
            return int(m.group(1))
        except Exception:
            pass
    return 0


# Официальный справочник профессорско-преподавательского состава 4 кафедр
TEACHERS_CATALOG = {
    # ─── Кафедра металлургических технологий и оборудования (МТиО) ───
    "нефёдов а.в.": {"full_name": "Нефёдов Андрей Викторович", "dept": "МТиО", "pos": "Заведующий кафедрой, к.п.н., доцент"},
    "нефедов а.в.": {"full_name": "Нефёдов Андрей Викторович", "dept": "МТиО", "pos": "Заведующий кафедрой, к.п.н., доцент"},
    "нефедов андрей викторович": {"full_name": "Нефёдов Андрей Викторович", "dept": "МТиО", "pos": "Заведующий кафедрой, к.п.н., доцент"},
    "нефёдов андрей викторович": {"full_name": "Нефёдов Андрей Викторович", "dept": "МТиО", "pos": "Заведующий кафедрой, к.п.н., доцент"},
    "братковский е.в.": {"full_name": "Братковский Евгений Владимирович", "dept": "МТиО", "pos": "Доцент кафедры, к.т.н."},
    "братковский евгений владимирович": {"full_name": "Братковский Евгений Владимирович", "dept": "МТиО", "pos": "Доцент кафедры, к.т.н."},
    "гавриш п.в.": {"full_name": "Гавриш Петр Владимирович", "dept": "МТиО", "pos": "Доцент кафедры, к.т.н."},
    "гавриш петр владимирович": {"full_name": "Гавриш Петр Владимирович", "dept": "МТиО", "pos": "Доцент кафедры, к.т.н."},
    "ганин д.р.": {"full_name": "Ганин Дмитрий Рудольфович", "dept": "МТиО", "pos": "Доцент кафедры, к.т.н."},
    "ганин дмитрий рудольфович": {"full_name": "Ганин Дмитрий Рудольфович", "dept": "МТиО", "pos": "Доцент кафедры, к.т.н."},
    "гриднева а.н.": {"full_name": "Гриднева Анастасия Николаевна", "dept": "МТиО", "pos": "Доцент кафедры"},
    "гриднева анастасия николаевна": {"full_name": "Гриднева Анастасия Николаевна", "dept": "МТиО", "pos": "Доцент кафедры"},
    "дема р.р.": {"full_name": "Дёма Роман Рафаэлевич", "dept": "МТиО", "pos": "Профессор кафедры, д.т.н."},
    "дёма р.р.": {"full_name": "Дёма Роман Рафаэлевич", "dept": "МТиО", "pos": "Профессор кафедры, д.т.н."},
    "дема роман рафаэлевич": {"full_name": "Дёма Роман Рафаэлевич", "dept": "МТиО", "pos": "Профессор кафедры, д.т.н."},
    "дёма роман рафаэлевич": {"full_name": "Дёма Роман Рафаэлевич", "dept": "МТиО", "pos": "Профессор кафедры, д.т.н."},
    "женин е.в.": {"full_name": "Женин Евгений Вячеславович", "dept": "МТиО", "pos": "Доцент кафедры"},
    "женин евгений вячеславович": {"full_name": "Женин Евгений Вячеславович", "dept": "МТиО", "pos": "Доцент кафедры"},
    "китанов а.а.": {"full_name": "Китанов Алексей Алексеевич", "dept": "МТиО", "pos": "Доцент кафедры"},
    "китанов алексей алексеевич": {"full_name": "Китанов Алексей Алексеевич", "dept": "МТиО", "pos": "Доцент кафедры"},
    "кузнецов м.с.": {"full_name": "Кузнецов Максим Сергеевич", "dept": "МТиО", "pos": "Доцент кафедры, к.т.н."},
    "кузнецов максим сергеевич": {"full_name": "Кузнецов Максим Сергеевич", "dept": "МТиО", "pos": "Доцент кафедры, к.т.н."},
    "куницина н.г.": {"full_name": "Куницина Наталья Геннадьевна", "dept": "МТиО", "pos": "Старший преподаватель"},
    "куницина наталья геннадьевна": {"full_name": "Куницина Наталья Геннадьевна", "dept": "МТиО", "pos": "Старший преподаватель"},
    "лактюшин а.а.": {"full_name": "Лактюшин Алексей Александрович", "dept": "МТиО", "pos": "Старший преподаватель"},
    "нефедьев с.п.": {"full_name": "Нефедьев Сергей Павлович", "dept": "МТиО", "pos": "Доцент кафедры, к.т.н."},
    "нефедьев сергей павлович": {"full_name": "Нефедьев Сергей Павлович", "dept": "МТиО", "pos": "Доцент кафедры, к.т.н."},
    "приступа а.н": {"full_name": "Приступа Анатолий Николаевич", "dept": "МТиО", "pos": "Доцент кафедры, к.т.н."},
    "приступа а.н.": {"full_name": "Приступа Анатолий Николаевич", "dept": "МТиО", "pos": "Доцент кафедры, к.т.н."},
    "степыко т.в.": {"full_name": "Степыко Татьяна Владимировна", "dept": "МТиО", "pos": "Старший преподаватель"},
    "степыко татьяна владимировна": {"full_name": "Степыко Татьяна Владимировна", "dept": "МТиО", "pos": "Старший преподаватель"},
    "табельская в.н.": {"full_name": "Табельская Вера Николаевна", "dept": "МТиО", "pos": "Старший преподаватель"},
    "табельская вера николаевна": {"full_name": "Табельская Вера Николаевна", "dept": "МТиО", "pos": "Старший преподаватель"},
    "фукс а.ю.": {"full_name": "Фукс Александр Юрьевич", "dept": "МТиО", "pos": "Доцент кафедры"},
    "фукс александр юрьевич": {"full_name": "Фукс Александр Юрьевич", "dept": "МТиО", "pos": "Доцент кафедры"},
    "харченко м.в.": {"full_name": "Харченко Максим Викторович", "dept": "МТиО", "pos": "Доцент кафедры, к.т.н."},
    "харченко максим викторович": {"full_name": "Харченко Максим Викторович", "dept": "МТиО", "pos": "Доцент кафедры, к.т.н."},
    "шаповалов а.н.": {"full_name": "Шаповалов Алексей Николаевич", "dept": "МТиО", "pos": "Профессор кафедры, д.т.н."},
    "шаповалов алексей николаевич": {"full_name": "Шаповалов Алексей Николаевич", "dept": "МТиО", "pos": "Профессор кафедры, д.т.н."},

    # ─── Кафедра математики и естествознания (МиЕ) ───
    "швалёва а.в.": {"full_name": "Швалёва Анна Викторовна", "dept": "МиЕ", "pos": "Заведующий кафедрой, к.п.н., доцент"},
    "швалева а.в.": {"full_name": "Швалёва Анна Викторовна", "dept": "МиЕ", "pos": "Заведующий кафедрой, к.п.н., доцент"},
    "швалёва анна викторовна": {"full_name": "Швалёва Анна Викторовна", "dept": "МиЕ", "pos": "Заведующий кафедрой, к.п.н., доцент"},
    "швалева анна викторовна": {"full_name": "Швалёва Анна Викторовна", "dept": "МиЕ", "pos": "Заведующий кафедрой, к.п.н., доцент"},
    "абдулвелеева р.р.": {"full_name": "Абдулвелеева Роза Рашитовна", "dept": "МиЕ", "pos": "Доцент кафедры, к.т.н."},
    "абдулвелеева роза рашитовна": {"full_name": "Абдулвелеева Роза Рашитовна", "dept": "МиЕ", "pos": "Доцент кафедры, к.т.н."},
    "алексеев д.и.": {"full_name": "Алексеев Данил Игоревич", "dept": "МиЕ", "pos": "Старший преподаватель"},
    "алексеев данил игоревич": {"full_name": "Алексеев Данил Игоревич", "dept": "МиЕ", "pos": "Старший преподаватель"},
    "бажуков д.о.": {"full_name": "Бажуков Дмитрий Олегович", "dept": "МиЕ", "pos": "Доцент кафедры, к.ф.-м.н."},
    "бажуков дмитрий олегович": {"full_name": "Бажуков Дмитрий Олегович", "dept": "МиЕ", "pos": "Доцент кафедры, к.ф.-м.н."},
    "белова м.н.": {"full_name": "Белова Марина Николаевна", "dept": "МиЕ", "pos": "Старший преподаватель"},
    "белова марина николаевна": {"full_name": "Белова Марина Николаевна", "dept": "МиЕ", "pos": "Старший преподаватель"},
    "варганова а.в.": {"full_name": "Варганова Александра Владимировна", "dept": "МиЕ", "pos": "Доцент кафедры, к.т.н."},
    "варганова александра владимировна": {"full_name": "Варганова Александра Владимировна", "dept": "МиЕ", "pos": "Доцент кафедры, к.т.н."},
    "косенок е.с.": {"full_name": "Косенок Евгений Сергеевич", "dept": "МиЕ", "pos": "Старший преподаватель"},
    "косенок е.c.": {"full_name": "Косенок Евгений Сергеевич", "dept": "МиЕ", "pos": "Старший преподаватель"},
    "косенок евгений сергеевич": {"full_name": "Косенок Евгений Сергеевич", "dept": "МиЕ", "pos": "Старший преподаватель"},
    "котов и.в.": {"full_name": "Котов Илья Вячеславович", "dept": "МиЕ", "pos": "Старший преподаватель"},
    "котов илья вячеславович": {"full_name": "Котов Илья Вячеславович", "dept": "МиЕ", "pos": "Старший преподаватель"},
    "леднов а.в.": {"full_name": "Леднов Анатолий Викторович", "dept": "МиЕ", "pos": "Старший преподаватель"},
    "леднов анатолий викторович": {"full_name": "Леднов Анатолий Викторович", "dept": "МиЕ", "pos": "Старший преподаватель"},
    "масалимов а.в.": {"full_name": "Масалимов Алексей Валерьевич", "dept": "МиЕ", "pos": "Старший преподаватель"},
    "масалимов алексей валерьевич": {"full_name": "Масалимов Алексей Валерьевич", "dept": "МиЕ", "pos": "Старший преподаватель"},
    "нефедова е.в.": {"full_name": "Нефедова Евгения Викторовна", "dept": "МиЕ", "pos": "Старший преподаватель"},
    "нефедова евгения викторовна": {"full_name": "Нефедова Евгения Викторовна", "dept": "МиЕ", "pos": "Старший преподаватель"},
    "новикова т.б.": {"full_name": "Новикова Татьяна Борисовна", "dept": "МиЕ", "pos": "Доцент кафедры, к.п.н."},
    "новикова татьяна борисовна": {"full_name": "Новикова Татьяна Борисовна", "dept": "МиЕ", "pos": "Доцент кафедры, к.п.н."},
    "рыбиков е.в.": {"full_name": "Рыбиков Евгений Владимирович", "dept": "МиЕ", "pos": "Доцент кафедры"},
    "рыбиков евгений владимирович": {"full_name": "Рыбиков Евгений Владимирович", "dept": "МиЕ", "pos": "Доцент кафедры"},
    "ткачева и.а.": {"full_name": "Ткачева Ирина Александровна", "dept": "МиЕ", "pos": "Старший преподаватель"},
    "ткачева ирина александровна": {"full_name": "Ткачева Ирина Александровна", "dept": "МиЕ", "pos": "Старший преподаватель"},
    "тобер с.г.": {"full_name": "Тобер Сергей Геннадьевич", "dept": "МиЕ", "pos": "Старший преподаватель"},
    "тобер сергей геннадьевич": {"full_name": "Тобер Сергей Геннадьевич", "dept": "МиЕ", "pos": "Старший преподаватель"},
    "филиппов е.г.": {"full_name": "Филиппов Евгений Георгиевич", "dept": "МиЕ", "pos": "Доцент кафедры, к.т.н."},
    "филиппов евгений георгиевич": {"full_name": "Филиппов Евгений Георгиевич", "dept": "МиЕ", "pos": "Доцент кафедры, к.т.н."},
    "филоненко т.п.": {"full_name": "Филоненко Татьяна Павловна", "dept": "МиЕ", "pos": "Старший преподаватель"},
    "филоненко татьяна павловна": {"full_name": "Филоненко Татьяна Павловна", "dept": "МиЕ", "pos": "Старший преподаватель"},
    "ячиков и.м.": {"full_name": "Ячиков Игорь Михайлович", "dept": "МиЕ", "pos": "Профессор кафедры, д.т.н."},
    "ячиков и.м": {"full_name": "Ячиков Игорь Михайлович", "dept": "МиЕ", "pos": "Профессор кафедры, д.т.н."},
    "ячиков игорь михайлович": {"full_name": "Ячиков Игорь Михайлович", "dept": "МиЕ", "pos": "Профессор кафедры, д.т.н."},

    # ─── Кафедра электроэнергетики и электротехники (ЭиЭ) ───
    "мажирина р.е.": {"full_name": "Мажирина Раиса Евгеньевна", "dept": "ЭиЭ", "pos": "Заведующий кафедрой, к.п.н., доцент"},
    "мажирина раиса евгеньевна": {"full_name": "Мажирина Раиса Евгеньевна", "dept": "ЭиЭ", "pos": "Заведующий кафедрой, к.п.н., доцент"},
    "антонов в.н.": {"full_name": "Антонов Вячеслав Николаевич", "dept": "ЭиЭ", "pos": "Доцент кафедры"},
    "антонов вячеслав николаевич": {"full_name": "Антонов Вячеслав Николаевич", "dept": "ЭиЭ", "pos": "Доцент кафедры"},
    "басков с.н.": {"full_name": "Басков Сергей Николаевич", "dept": "ЭиЭ", "pos": "Доцент кафедры, к.т.н."},
    "басков сергей николаевич": {"full_name": "Басков Сергей Николаевич", "dept": "ЭиЭ", "pos": "Доцент кафедры, к.т.н."},
    "белых д.в.": {"full_name": "Белых Дарья Васильевна", "dept": "ЭиЭ", "pos": "Старший преподаватель"},
    "белых дарья васильевна": {"full_name": "Белых Дарья Васильевна", "dept": "ЭиЭ", "pos": "Старший преподаватель"},
    "бушуев а.н.": {"full_name": "Бушуев Антон Николаевич", "dept": "ЭиЭ", "pos": "Доцент кафедры"},
    "бушуев антон николаевич": {"full_name": "Бушуев Антон Николаевич", "dept": "ЭиЭ", "pos": "Доцент кафедры"},
    "лицин к.в.": {"full_name": "Лицин Константин Владимирович", "dept": "ЭиЭ", "pos": "Доцент кафедры, к.т.н."},
    "лицин константин владимирович": {"full_name": "Лицин Константин Владимирович", "dept": "ЭиЭ", "pos": "Доцент кафедры, к.т.н."},

    # ─── Кафедра гуманитарных и социально-экономических наук (ГиСЭН) ───
    "измайлова а.с.": {"full_name": "Измайлова Анна Сергеевна", "dept": "ГиСЭН", "pos": "Заведующий кафедрой, к.э.н., доцент"},
    "измайлова анна сергеевна": {"full_name": "Измайлова Анна Сергеевна", "dept": "ГиСЭН", "pos": "Заведующий кафедрой, к.э.н., доцент"},
    "болдырева н.п.": {"full_name": "Болдырева Нина Павловна", "dept": "ГиСЭН", "pos": "Доцент кафедры"},
    "болдырева нина павловна": {"full_name": "Болдырева Нина Павловна", "dept": "ГиСЭН", "pos": "Доцент кафедры"},
    "выскребцев б.с.": {"full_name": "Выскребцев Богдан Сергеевич", "dept": "ГиСЭН", "pos": "Преподаватель"},
    "выскребцев богдан сергеевич": {"full_name": "Выскребцев Богдан Сергеевич", "dept": "ГиСЭН", "pos": "Преподаватель"},
    "елисеева и.а.": {"full_name": "Елисеева Ирина Александровна", "dept": "ГиСЭН", "pos": "Доцент кафедры, к.филол.н."},
    "елисеева ирина александровна": {"full_name": "Елисеева Ирина Александровна", "dept": "ГиСЭН", "pos": "Доцент кафедры, к.филол.н."},
    "зенченко и.в.": {"full_name": "Зенченко Ирина Владимировна", "dept": "ГиСЭН", "pos": "Доцент кафедры, к.э.н."},
    "зенченко ирина владимировна": {"full_name": "Зенченко Ирина Владимировна", "dept": "ГиСЭН", "pos": "Доцент кафедры, к.э.н."},
    "киселева о.в.": {"full_name": "Киселева Ольга Вячеславовна", "dept": "ГиСЭН", "pos": "Доцент кафедры, к.и.н."},
    "киселева ольга вячеславовна": {"full_name": "Киселева Ольга Вячеславовна", "dept": "ГиСЭН", "pos": "Доцент кафедры, к.и.н."},
    "мусафиров м.к.": {"full_name": "Мусафиров Михаил Константинович", "dept": "ГиСЭН", "pos": "Доцент кафедры, к.и.н."},
    "мусафиров михаил константинович": {"full_name": "Мусафиров Михаил Константинович", "dept": "ГиСЭН", "pos": "Доцент кафедры, к.и.н."},
    "нечетов в.г.": {"full_name": "Нечетов Василий Геннадьевич", "dept": "ГиСЭН", "pos": "Доцент кафедры, к.э.н."},
    "нечетов василий геннадьевич": {"full_name": "Нечетов Василий Геннадьевич", "dept": "ГиСЭН", "pos": "Доцент кафедры, к.э.н."},
    "свечникова в.в.": {"full_name": "Свечникова Виктория Владимировна", "dept": "ГиСЭН", "pos": "Старший преподаватель"},
    "свечникова виктория владимировна": {"full_name": "Свечникова Виктория Владимировна", "dept": "ГиСЭН", "pos": "Старший преподаватель"},
    "торшина а.в.": {"full_name": "Торшина Анна Вячеславовна", "dept": "ГиСЭН", "pos": "Старший преподаватель"},
    "торшина анна вячеславовна": {"full_name": "Торшина Анна Вячеславовна", "dept": "ГиСЭН", "pos": "Старший преподаватель"},
}


class Command(BaseCommand):
    help = "Seed database with real semester workload from Excel and survey templates (Autumn / Spring)"

    @transaction.atomic
    def handle(self, *args, **options):
        self.stdout.write("Очистка устаревших данных...")

        # 0. Очистка зависимых сущностей в правильном порядке
        Grade.objects.all().delete()
        Student.objects.all().delete()
        SurveyAnswer.objects.all().delete()
        SurveyAssignment.objects.all().delete()
        TeacherRecommendation.objects.all().delete()
        SurveyQuestion.objects.all().delete()
        SurveyTemplate.objects.all().delete()
        Workload.objects.all().delete()
        StudyGroup.objects.all().delete()
        Discipline.objects.all().delete()
        Teacher.objects.all().delete()
        Department.objects.all().delete()
        User.objects.filter(username__startswith="t_").delete()
        AuditLog.objects.all().delete()

        try:
            cache.clear()
        except Exception:
            pass

        self.stdout.write("1. Создание 4 официальных кафедр НФ НИТУ МИСИС...")
        dept_mtio, _ = Department.objects.get_or_create(
            code="МТиО",
            defaults={"name": "Кафедра металлургических технологий и оборудования"},
        )
        dept_mie, _ = Department.objects.get_or_create(
            code="МиЕ",
            defaults={"name": "Кафедра математики и естествознания"},
        )
        dept_eie, _ = Department.objects.get_or_create(
            code="ЭиЭ",
            defaults={"name": "Кафедра электроэнергетики и электротехники"},
        )
        dept_gisen, _ = Department.objects.get_or_create(
            code="ГиСЭН",
            defaults={"name": "Кафедра гуманитарных и социально-экономических наук"},
        )

        dept_map = {
            "МТиО": dept_mtio,
            "МиЕ": dept_mie,
            "ЭиЭ": dept_eie,
            "ГиСЭН": dept_gisen,
        }

        self.stdout.write("2. Создание ключевых пользователей системы...")
        # Зав. кафедрой ГиСЭН: Измайлова Анна Сергеевна
        head_user, _ = User.objects.get_or_create(
            username="head",
            defaults={
                "role": "head",
                "first_name": "Анна",
                "last_name": "Измайлова",
                "email": "izmailova@misis.ru",
                "is_staff": True,
            },
        )
        head_user.first_name = "Анна"
        head_user.last_name = "Измайлова"
        head_user.email = "izmailova@misis.ru"
        head_user.role = "head"
        head_user.is_staff = True
        head_user.set_password("head12345")
        head_user.save()

        # Преподаватель 1: Нечетов Василий Геннадьевич (кафедра ГиСЭН)
        t1_user, _ = User.objects.get_or_create(
            username="teacher1",
            defaults={
                "role": "teacher",
                "first_name": "Василий",
                "last_name": "Нечетов",
                "email": "nechetov@misis.ru",
            },
        )
        t1_user.first_name = "Василий"
        t1_user.last_name = "Нечетов"
        t1_user.email = "nechetov@misis.ru"
        t1_user.role = "teacher"
        t1_user.set_password("teacher12345")
        t1_user.save()

        # Преподаватель 2: Нефедова Евгения Викторовна (кафедра МиЕ)
        t2_user, _ = User.objects.get_or_create(
            username="teacher2",
            defaults={
                "role": "teacher",
                "first_name": "Евгения",
                "last_name": "Нефедова",
                "email": "nefedova@misis.ru",
            },
        )
        t2_user.first_name = "Евгения"
        t2_user.last_name = "Нефедова"
        t2_user.email = "nefedova@misis.ru"
        t2_user.role = "teacher"
        t2_user.set_password("teacher22345")
        t2_user.save()

        # Администратор
        admin_user, _ = User.objects.get_or_create(
            username="admin",
            defaults={
                "role": "admin",
                "first_name": "Системный",
                "last_name": "Администратор",
                "email": "admin@misis.ru",
                "is_staff": True,
                "is_superuser": True,
            },
        )
        admin_user.role = "admin"
        admin_user.is_staff = True
        admin_user.is_superuser = True
        admin_user.set_password("admin12345")
        admin_user.save()

        # Профиль заведующего кафедрой (Измайлова Анна Сергеевна)
        teacher_head, _ = Teacher.objects.get_or_create(
            user=head_user,
            defaults={
                "full_name": "Измайлова Анна Сергеевна",
                "department": dept_gisen,
                "position": "Заведующий кафедрой, к.э.н., доцент",
                "hours_limit": 900.0,
            },
        )
        teacher_head.full_name = "Измайлова Анна Сергеевна"
        teacher_head.department = dept_gisen
        teacher_head.position = "Заведующий кафедрой, к.э.н., доцент"
        teacher_head.hours_limit = 900.0
        teacher_head.save()

        teachers_cache = {
            "измайлова а.с.": teacher_head,
            "измайлова анна сергеевна": teacher_head,
        }

        # Преподаватель Нечетов
        teacher_nechetov, _ = Teacher.objects.get_or_create(
            user=t1_user,
            defaults={
                "full_name": "Нечетов Василий Геннадьевич",
                "department": dept_gisen,
                "position": "Доцент кафедры, к.э.н.",
                "hours_limit": 850.0,
            },
        )
        teacher_nechetov.full_name = "Нечетов Василий Геннадьевич"
        teacher_nechetov.department = dept_gisen
        teacher_nechetov.position = "Доцент кафедры, к.э.н."
        teacher_nechetov.save()
        teachers_cache["нечетов в.г."] = teacher_nechetov
        teachers_cache["нечетов василий геннадьевич"] = teacher_nechetov

        # Преподаватель Нефедова
        teacher_nefedova, _ = Teacher.objects.get_or_create(
            user=t2_user,
            defaults={
                "full_name": "Нефедова Евгения Викторовна",
                "department": dept_mie,
                "position": "Старший преподаватель",
                "hours_limit": 800.0,
            },
        )
        teacher_nefedova.full_name = "Нефедова Евгения Викторовна"
        teacher_nefedova.department = dept_mie
        teacher_nefedova.position = "Старший преподаватель"
        teacher_nefedova.save()
        teachers_cache["нефедова е.в."] = teacher_nefedova
        teachers_cache["нефедова евгения викторовна"] = teacher_nefedova

        def get_or_create_teacher(name_str):
            cleaned = clean_teacher_name(name_str)
            if not cleaned or cleaned.lower() == "вакансия":
                return teacher_head

            key = cleaned.lower()
            if key in teachers_cache:
                return teachers_cache[key]

            # Ищем в каталоге
            meta = TEACHERS_CATALOG.get(key)
            if not meta:
                # Попробуем без точек
                for k, v in TEACHERS_CATALOG.items():
                    if k.replace(".", "") == key.replace(".", ""):
                        meta = v
                        break

            if meta:
                full_name = meta["full_name"]
                dept_obj = dept_map.get(meta["dept"], dept_gisen)
                pos = meta["pos"]
            else:
                full_name = cleaned
                dept_obj = dept_gisen
                pos = "Доцент кафедры"

            if full_name == "Измайлова Анна Сергеевна":
                teachers_cache[key] = teacher_head
                return teacher_head

            # Проверяем, существует ли уже преподаватель с таким ФИО
            existing = Teacher.objects.filter(full_name=full_name).first()
            if existing:
                existing.department = dept_obj
                existing.position = pos
                existing.save()
                teachers_cache[key] = existing
                return existing

            username_clean = f"t_{abs(hash(full_name)) % 100000}"
            u, _ = User.objects.get_or_create(
                username=username_clean,
                defaults={
                    "role": "teacher",
                    "first_name": full_name.split()[1] if len(full_name.split()) > 1 else full_name,
                    "last_name": full_name.split()[0],
                    "email": f"{username_clean}@misis.ru",
                },
            )
            u.set_password("teacher12345")
            u.save()

            tch = Teacher.objects.create(
                user=u,
                full_name=full_name,
                department=dept_obj,
                position=pos,
                hours_limit=850.0,
            )
            teachers_cache[key] = tch
            teachers_cache[full_name.lower()] = tch
            return tch

        self.stdout.write("3. Создание 20 академических учебных групп (1-4 курсы)...")
        # 4 курса, по 5 групп на курс
        # БПИ: 09.03.01, БХТ: 18.03.01, БМТ: 22.03.02, БЭЭ: 13.03.02, БТМО: 15.03.02
        group_specs = [
            ("БПИ", "09.03.01", "Прикладная информатика"),
            ("БХТ", "18.03.01", "Химическая технология"),
            ("БМТ", "22.03.02", "Металлургия"),
            ("БЭЭ", "13.03.02", "Электроэнергетика и электротехника"),
            ("БТМО", "15.03.02", "Технологические машины и оборудование"),
        ]

        study_groups = {}
        for course in [1, 2, 3, 4]:
            year = 2027 - course  # 1к -> 2026, 2к -> 2025, 3к -> 2024, 4к -> 2023
            suffix = str(year)[-2:]
            for prefix, d_code, d_name in group_specs:
                grp_name = f"{prefix}-{suffix}"
                grp, _ = StudyGroup.objects.get_or_create(
                    name=grp_name,
                    defaults={
                        "course": course,
                        "year": year,
                        "direction_code": d_code,
                        "direction_name": d_name,
                    },
                )
                grp.course = course
                grp.year = year
                grp.direction_code = d_code
                grp.direction_name = d_name
                grp.save()
                study_groups[grp_name] = grp

        self.stdout.write(f"Создано {len(study_groups)} учебных групп.")

        self.stdout.write("4. Создание студентов...")
        first_names = ["Максим", "Анна", "Денис", "Артём", "Ольга", "Илья", "Дмитрий", "Дарья", "Полина", "Никита"]
        last_names = ["Смирнов", "Кузнецова", "Васильев", "Соколов", "Ковалева", "Новиков", "Волков", "Морозова", "Лебедева", "Попов"]
        students_all = []
        for grp_name, grp in study_groups.items():
            for i in range(5):
                fn = first_names[(hash(grp_name) + i) % len(first_names)]
                ln = last_names[(hash(grp_name) + i * 2) % len(last_names)]
                st = Student.objects.create(
                    full_name=f"{ln} {fn}",
                    email=f"{grp_name.lower()}_{i+1}@misis.ru",
                    group=grp,
                )
                students_all.append(st)

        self.stdout.write(f"Создано {len(students_all)} студентов.")

        self.stdout.write("5. Создание шаблонов опроса для Осеннего и Весеннего семестров...")
        template_autumn, _ = SurveyTemplate.objects.get_or_create(
            semester="2026-1",
            academic_year="2026-2027",
            defaults={
                "title": "Оценка качества преподавания (Осенний семестр 2026/2027)",
                "description": "Анонимный опрос студентов филиала по итогам осеннего семестра",
                "is_active": True,
            },
        )
        template_autumn.title = "Оценка качества преподавания (Осенний семестр 2026/2027)"
        template_autumn.is_active = True
        template_autumn.save()

        template_spring, _ = SurveyTemplate.objects.get_or_create(
            semester="2026-2",
            academic_year="2026-2027",
            defaults={
                "title": "Оценка качества преподавания (Весенний семестр 2026/2027)",
                "description": "Анонимный опрос студентов филиала по итогам весеннего семестра",
                "is_active": True,
            },
        )
        template_spring.title = "Оценка качества преподавания (Весенний семестр 2026/2027)"
        template_spring.is_active = True
        template_spring.save()

        # 12 академических критериев оценивания (привязаны к 5 блокам качества)
        questions_def = [
            ("relevance", "Блок 1. Организационно-методическое обеспечение дисциплины (промежуточная аттестация, дедлайны, критерии)", "scale_5", [], 1),
            ("clarity", "Блок 2. Качество проведения занятий и подача материала (структурированность лекций, разбор практики)", "scale_5", [], 2),
            ("fairness", "Блок 3. Выполнение учебного плана и дисциплина проведения занятий (соблюдение расписания, консультации)", "scale_5", [], 3),
            ("ethics", "Блок 4. Формирование атмосферы вовлечённости в изучение дисциплины (педагогический такт, интерес, этикет)", "scale_5", [], 4),
            ("facilities", "Блок 5. Наполнение и использование электронного курса в LMS Moodle (актуальные материалы, обратная связь)", "scale_5", [], 5),
            ("general", "Замечания и дефекты: чего вам не хватило на занятиях и в организации дисциплины", "text", [], 6),
            ("general", "Благодарность преподавателю: что вам больше всего понравилось в работе и содержании курса", "text", [], 7),
            ("general", "Регулярность обратной связи преподавателя по выполненным заданиям и проектам", "single_choice", [
                "На каждом занятии или в течение 3 дней после сдачи",
                "Раз в 1-2 недели по установленному графику",
                "Только перед контрольными точками и зачетом",
                "Обратная связь практически отсутствовала"
            ], 8),
            ("general", "Степень использования современных цифровых сервисов и LMS в процессе обучения", "single_choice", [
                "Постоянно (LMS Moodle, электронные тесты, интерактивные презентации)",
                "Периодически по ключевым сложным темам",
                "Редко, преимущественно классический лекционный формат",
                "Не использовались"
            ], 9),
            ("general", "Оцените общий темп освоения учебного материала и баланс учебной нагрузки", "single_choice", [
                "Оптимальный темп, нагрузка распределена равномерно",
                "Интенсивный темп, требовалось много самостоятельной работы",
                "Умеренный темп, материал давался легко",
                "Перегруженный темп, времени на качественное освоение не хватало"
            ], 10),
            ("relevance", "Какие формы работы на занятиях оказались для вас наиболее эффективными и полезными?", "multiple_choice", [
                "Разбор реальных отраслевых и производственных кейсов",
                "Интерактивные дискуссии и семинарские обсуждения",
                "Практические расчетные задания и моделирование",
                "Работа в малых командах над проектом",
                "Индивидуальные консультации с разбором ошибок"
            ], 11),
            ("general", "Ваши конкретные предложения преподавателю по развитию курса", "text", [], 12),
        ]

        questions_autumn = []
        for cat, txt, qtype, opts, ord_idx in questions_def:
            q = SurveyQuestion.objects.create(
                template=template_autumn,
                text=txt,
                category=cat,
                question_type=qtype,
                options=opts,
                order=ord_idx,
            )
            questions_autumn.append(q)

        questions_spring = []
        for cat, txt, qtype, opts, ord_idx in questions_def:
            q = SurveyQuestion.objects.create(
                template=template_spring,
                text=txt,
                category=cat,
                question_type=qtype,
                options=opts,
                order=ord_idx,
            )
            questions_spring.append(q)

        # Поиск файлов нагрузки
        repo_root = Path(__file__).resolve().parent.parent.parent.parent.parent
        base_candidates = [
            repo_root / "data" / "workload_input",
            repo_root / "backend" / "data" / "workload_input",
            Path("/app/data/workload_input"),
            Path("/app/backend/data/workload_input"),
            Path.cwd() / "data" / "workload_input",
            Path.cwd() / "backend" / "data" / "workload_input",
            Path("backend/data/workload_input"),
            Path("data/workload_input"),
            Path("."),
        ]

        excel_files = []
        for course_num in [1, 2, 3, 4]:
            fname = f"SEMESTROVAYa_NAGRUZKA_{course_num}k.xlsx"
            found_path = None
            for cand in base_candidates:
                p = cand / fname
                if p.exists():
                    found_path = p
                    break
            if found_path:
                excel_files.append((found_path, course_num))
            else:
                self.stdout.write(self.style.WARNING(f"Файл {fname} не найден в кандидатах"))

        self.stdout.write(f"Найдено файлов для импорта: {len(excel_files)}")

        disciplines_cache = {}
        disc_counter = 1

        def get_or_create_discipline(name_str, lec_h, pz_h, lab_h):
            nonlocal disc_counter
            name_clean = str(name_str).strip()
            if name_clean in disciplines_cache:
                return disciplines_cache[name_clean]

            lesson_t = "lecture"
            if lab_h > 0:
                lesson_t = "lab"
            elif pz_h > 0:
                lesson_t = "practice"

            tot_h = (lec_h + pz_h + lab_h) or 72
            disc_code = f"DISC-{disc_counter:03d}"
            disc_counter += 1

            disc, _ = Discipline.objects.get_or_create(
                name=name_clean,
                defaults={
                    "code": disc_code,
                    "total_hours": tot_h,
                    "lesson_type": lesson_t,
                },
            )
            disciplines_cache[name_clean] = disc
            return disc

        workloads_created = 0
        assignments_autumn = []
        assignments_spring = []

        day_rotator = 1
        lesson_rotator = 1
        rooms_list = ["101", "203", "204", "312", "315", "401", "408", "Акт. зал", "Комп. класс 1", "Комп. класс 2"]

        for fpath, course_num in excel_files:
            wb = openpyxl.load_workbook(fpath, data_only=True)
            for sname in wb.sheetnames:
                sheet = wb[sname]
                grp_obj = study_groups.get(sname.strip())
                if not grp_obj:
                    continue

                current_semester = None
                for r in range(1, sheet.max_row + 1):
                    row_str = " ".join([str(sheet.cell(r, c).value or "") for c in range(1, 12)])
                    sem_match = re.search(r"СЕМЕСТР\s*(\d+)", row_str, re.IGNORECASE)
                    if sem_match:
                        current_semester = int(sem_match.group(1))
                        continue

                    if current_semester is None:
                        continue

                    cell_b = sheet.cell(r, 2).value
                    cell_c = sheet.cell(r, 3).value
                    if not cell_b:
                        continue
                    disc_name = str(cell_b).strip()
                    if disc_name.lower() in ["дисциплина", "итог", "всего"]:
                        continue
                    if any(k in disc_name.lower() for k in ["учебный год", "группа", "зав. кафедрой", "декан", "руководитель", "подпись"]):
                        continue

                    lec_h = parse_hours(sheet.cell(r, 8).value)
                    pz_h = parse_hours(sheet.cell(r, 9).value)
                    lab_h = parse_hours(sheet.cell(r, 10).value)

                    disc_obj = get_or_create_discipline(disc_name, lec_h, pz_h, lab_h)

                    # Преподаватели
                    teacher_raw = str(cell_c).strip() if cell_c else ""
                    row_teachers = []
                    if teacher_raw and teacher_raw != "None":
                        for t_str in re.split(r"[\n\r/]+", teacher_raw):
                            c_name = clean_teacher_name(t_str)
                            if c_name and not c_name.isdigit() and len(c_name) > 2 and "преподаватель" not in c_name.lower():
                                row_teachers.append(get_or_create_teacher(c_name))

                    if not row_teachers:
                        row_teachers = [teacher_head]

                    is_autumn = (current_semester % 2 == 1)
                    sem_code = "2026-1" if is_autumn else "2026-2"
                    target_template = template_autumn if is_autumn else template_spring

                    for tch in row_teachers:
                        hours_p = (lec_h + pz_h + lab_h) or 36
                        hours_f = max(0, int(hours_p * 0.75))

                        wl = Workload.objects.create(
                            teacher=tch,
                            discipline=disc_obj,
                            group=grp_obj,
                            semester=sem_code,
                            hours_plan=hours_p,
                            hours_fact=hours_f,
                            day_of_week=(day_rotator % 6) + 1,
                            lesson_number=(lesson_rotator % 5) + 1,
                            room=rooms_list[(day_rotator + lesson_rotator) % len(rooms_list)],
                        )
                        workloads_created += 1
                        day_rotator += 1
                        lesson_rotator += 1

                        # Назначение для опроса студентов (привязываем к реальной кафедре преподавателя!)
                        assign = SurveyAssignment.objects.create(
                            template=target_template,
                            teacher=tch,
                            discipline=disc_obj,
                            group=grp_obj,
                            department=tch.department or dept_gisen,
                            is_open=True,
                        )
                        if is_autumn:
                            assignments_autumn.append(assign)
                        else:
                            assignments_spring.append(assign)

        self.stdout.write(f"Создано строк нагрузки: {workloads_created}")
        self.stdout.write(f"Создано назначений опроса (Осень): {len(assignments_autumn)}")
        self.stdout.write(f"Создано назначений опроса (Весна): {len(assignments_spring)}")

        self.stdout.write("7. Генерация успеваемости студентов (Grades)...")
        grade_count = 0
        for grp in study_groups.values():
            grp_students = Student.objects.filter(group=grp)
            grp_workloads = Workload.objects.filter(group=grp)[:6]
            for st in grp_students:
                for wl in grp_workloads:
                    # Разброс оценок: от 3.0 до 5.0
                    grade_num = [4.0, 5.0, 3.0, 5.0, 4.0, 5.0][(abs(hash(st.full_name) + hash(wl.discipline.name))) % 6]
                    Grade.objects.create(
                        student=st,
                        discipline=wl.discipline,
                        teacher=wl.teacher,
                        semester=wl.semester,
                        grade=grade_num,
                        source=Grade.MOODLE,
                        date=date(2026, 10, 1),
                    )
                    grade_count += 1
        self.stdout.write(f"Создано {grade_count} записей успеваемости.")

        self.stdout.write("8. Заполнение реалистичных ответов студентов (Survey Answers)...")

        # Отзывы для Измайловой Анны Сергеевны
        head_feedbacks = [
            ([5, 5, 5, 5, 5], "Анна Сергеевна прекрасно объясняет методологию ведения IT-проектов, всегда готова проконсультировать по сложным практическим задачам!", "Разбор реальных отраслевых и производственных кейсов"),
            ([5, 5, 4, 5, 5], "Один из лучших и самых полезных курсов! Разобрали реальные бизнес-кейсы по планированию ресурсов и Scrum-процессам.", "Работа в малых командах над проектом"),
            ([5, 5, 5, 5, 5], "Преподаватель очень доброжелателен и справедлив. Четкие критерии, полезные дедлайны и интересные дискуссии.", "Интерактивные дискуссии и семинарские обсуждения"),
            ([5, 4, 5, 5, 4], "Курс дал мощное понимание экономической эффективности и управления рисками в проектах.", "Практические расчетные задания и моделирование"),
        ]

        sample_feedback_autumn = [
            ([5, 5, 4, 5, 4], "Лекции читаются очень структурированно и с живыми практическими примерами.", "Разбор реальных отраслевых и производственных кейсов"),
            ([4, 4, 4, 5, 5], "Доброжелательный преподаватель, оперативная обратная связь по лабораторным.", "Интерактивные дискуссии и семинарские обсуждения"),
            ([5, 4, 5, 4, 4], "Отличный курс! Разобрали современные методики и требования индустрии.", "Практические расчетные задания и моделирование"),
            ([4, 4, 4, 4, 4], "Интересный материал, понятные дескрипторы оценивания.", "Индивидуальные консультации с разбором ошибок"),
            ([5, 5, 5, 5, 5], "Все четко, логично, вовремя и справедливо.", "Разбор реальных отраслевых и производственных кейсов"),
        ]

        sample_feedback_spring = [
            ([5, 4, 5, 5, 4], "Весенний блок дисциплины еще интереснее, много практических задач.", "Разбор реальных отраслевых и производственных кейсов"),
            ([4, 4, 4, 4, 4], "Понятные дескрипторы оценивания, баллы в БРС выставляются вовремя.", "Практические расчетные задания и моделирование"),
            ([5, 5, 4, 5, 5], "Качественные материалы в Moodle, презентации всегда доступны онлайн.", "Интерактивные дискуссии и семинарские обсуждения"),
            ([4, 4, 5, 4, 4], "Преподаватель мотивирует изучать глубже и участвовать в конференциях.", "Работа в малых командах над проектом"),
        ]

        def seed_answers_for_assignment(a, q_list, feedback_pool, num_students=3):
            count = 0
            for idx in range(num_students):
                scores, comment_pos, format_choice = feedback_pool[(hash(a.teacher.full_name) + idx) % len(feedback_pool)]
                s_hash = uuid.uuid4().hex
                # 5 матричных шкал (1-5)
                for q_idx in range(5):
                    SurveyAnswer.objects.create(
                        assignment=a,
                        question=q_list[q_idx],
                        score=scores[q_idx],
                        submission_hash=s_hash,
                    )
                # Вопрос 6: замечания
                SurveyAnswer.objects.create(
                    assignment=a,
                    question=q_list[5],
                    text_response="" if scores[0] >= 4 else "Хотелось бы больше времени на практические расчеты.",
                    submission_hash=s_hash,
                )
                # Вопрос 7: благодарность
                SurveyAnswer.objects.create(
                    assignment=a,
                    question=q_list[6],
                    text_response=comment_pos,
                    submission_hash=s_hash,
                )
                # Вопрос 8: обратная связь
                SurveyAnswer.objects.create(
                    assignment=a,
                    question=q_list[7],
                    text_response="На каждом занятии или в течение 3 дней после сдачи",
                    submission_hash=s_hash,
                )
                # Вопрос 9: LMS
                SurveyAnswer.objects.create(
                    assignment=a,
                    question=q_list[8],
                    text_response="Постоянно (LMS Moodle, электронные тесты, интерактивные презентации)",
                    submission_hash=s_hash,
                )
                # Вопрос 10: темп
                SurveyAnswer.objects.create(
                    assignment=a,
                    question=q_list[9],
                    text_response="Оптимальный темп, нагрузка распределена равномерно",
                    submission_hash=s_hash,
                )
                # Вопрос 11: полезные формы
                SurveyAnswer.objects.create(
                    assignment=a,
                    question=q_list[10],
                    text_response=format_choice,
                    submission_hash=s_hash,
                )
                # Вопрос 12: пожелания
                SurveyAnswer.objects.create(
                    assignment=a,
                    question=q_list[11],
                    text_response="Продолжать развитие курса и практических отраслевых кейсов.",
                    submission_hash=s_hash,
                )
                count += 1
            return count

        total_sessions = 0

        # 1. Приоритетно сидируем все назначения Измайловой Анны Сергеевны!
        head_assigns = [a for a in assignments_autumn + assignments_spring if a.teacher == teacher_head]
        self.stdout.write(f"Найдено назначений Измайловой А.С.: {len(head_assigns)}")
        for a in head_assigns:
            q_list = questions_autumn if a.template.semester == "2026-1" else questions_spring
            total_sessions += seed_answers_for_assignment(a, q_list, head_feedbacks, num_students=4)

        # 2. Сидируем назначения по каждой кафедре (по 10-15 назначений на кафедру)
        for d in [dept_gisen, dept_mie, dept_mtio, dept_eie]:
            dept_assigns_autumn = [a for a in assignments_autumn if a.department == d and a.teacher != teacher_head][:12]
            dept_assigns_spring = [a for a in assignments_spring if a.department == d and a.teacher != teacher_head][:12]
            for a in dept_assigns_autumn:
                total_sessions += seed_answers_for_assignment(a, questions_autumn, sample_feedback_autumn, num_students=3)
            for a in dept_assigns_spring:
                total_sessions += seed_answers_for_assignment(a, questions_spring, sample_feedback_spring, num_students=3)

        self.stdout.write(f"Всего сгенерировано сессий опросов: {total_sessions}")

        # Создаем официальную рекомендацию для Измайловой А.С.
        TeacherRecommendation.objects.create(
            teacher=teacher_head,
            category="relevance",
            recommendation_text=(
                "Высокий уровень студенческой удовлетворенности (95%+). "
                "Рекомендуется транслировать опыт практико-ориентированного проектного обучения "
                "и применения кейс-метода на другие кафедры филиала."
            ),
            semester="2026-1",
            source="head",
            status="reviewed",
        )

        self.stdout.write(self.style.SUCCESS("Импорт реальной нагрузки по 4 кафедрам и сидирование успешно завершены!"))
